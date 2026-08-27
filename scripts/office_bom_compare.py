"""Compare two BOM files (.xls/.xlsx) keyed by designator column.

usage: office_bom_compare.py <a> <b> [key-hint]
读出两表（xlrd 处理 .xls / openpyxl 处理 .xlsx），逐文件扫描前若干行定位「含位号表头
（位号/refdes/designator/reference/编号/料号，表头先去点去空格归一化，或 key-hint 命中）」
的表头行与位号列（缺省第 0 列），按位号切分键并输出 公共/仅 A/仅 B/变更 四类差异（样例有界）。
"""

import json
import re
import sys

sys.stdout.reconfigure(encoding="utf-8")

MAX_SAMPLE = 30
KEY_HEADER_RE = re.compile(r"位号|refdes|designator|reference|编号|料号|品号|物料编码")


def read_rows(path):
    ext = path.lower().rsplit(".", 1)[-1] if "." in path else ""
    if ext == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([normalize_cell(c) for c in row])
            if len(rows) >= 2000:
                break
        wb.close()
        return rows
    import xlrd

    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(min(sheet.nrows, 2000)):
        rows.append([normalize_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)])
    return rows


def normalize_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def norm_header(value):
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]", "", value.lower())


def split_keys(cell):
    return sorted({k for k in re.split(r"[,，、;；/\s]+", cell) if k})


def find_header(rows, hint):
    """返回 (header_row_index, key_col)；找不到位号表头时退回首行/第 0 列。"""
    limit = min(len(rows), 10)
    if hint:
        low = norm_header(hint)
        for r in range(limit):
            for c, h in enumerate(rows[r]):
                if low and low in norm_header(h):
                    return r, c
    for r in range(limit):
        for c, h in enumerate(rows[r]):
            if KEY_HEADER_RE.search(norm_header(h)):
                return r, c
    return 0, 0


def build_map(rows, header_index, key_col):
    key_map = {}
    for row in rows[header_index + 1 :]:
        if key_col >= len(row):
            continue
        for key in split_keys(row[key_col]):
            key_map.setdefault(key.upper(), []).append(row)
    return key_map


def diff_fields(row_a, row_b, header, key_col):
    width = max(len(row_a), len(row_b))
    diffs = []
    for i in range(width):
        if i == key_col:
            continue
        a = row_a[i] if i < len(row_a) else ""
        b = row_b[i] if i < len(row_b) else ""
        if a != b:
            name = header[i] if i < len(header) else f"列{i+1}"
            diffs.append({"field": name, "before": a, "after": b})
    return diffs


def main() -> None:
    if len(sys.argv) < 3:
        print(json.dumps({"ok": False, "error": "usage: office_bom_compare.py <a> <b> [key-hint]"}))
        sys.exit(1)
    path_a, path_b = sys.argv[1], sys.argv[2]
    hint = sys.argv[3] if len(sys.argv) > 3 else ""
    try:
        rows_a = read_rows(path_a)
        rows_b = read_rows(path_b)
        header_a_idx, key_col_a = find_header(rows_a, hint)
        header_b_idx, key_col_b = find_header(rows_b, hint)
        map_a = build_map(rows_a, header_a_idx, key_col_a)
        map_b = build_map(rows_b, header_b_idx, key_col_b)
        header_a = rows_a[header_a_idx] if rows_a else []
        header_b = rows_b[header_b_idx] if rows_b else []
        keys_a = set(map_a)
        keys_b = set(map_b)
        common = sorted(keys_a & keys_b)
        only_a = sorted(keys_a - keys_b)
        only_b = sorted(keys_b - keys_a)
        changed = []
        for key in common:
            diffs = diff_fields(map_a[key][0], map_b[key][0], header_a, key_col_a)
            if diffs:
                changed.append({"key": key, "fields": diffs})
        print(json.dumps({
            "ok": True,
            "file_a": path_a,
            "file_b": path_b,
            "key_column_a": header_a[key_col_a] if key_col_a < len(header_a) else f"列{key_col_a+1}",
            "key_column_b": header_b[key_col_b] if key_col_b < len(header_b) else f"列{key_col_b+1}",
            "rows_a": len(rows_a) - header_a_idx - 1,
            "rows_b": len(rows_b) - header_b_idx - 1,
            "common": len(common),
            "only_a_count": len(only_a),
            "only_b_count": len(only_b),
            "changed_count": len(changed),
            "only_in_a": only_a[:MAX_SAMPLE],
            "only_in_b": only_b[:MAX_SAMPLE],
            "changed": changed[:MAX_SAMPLE],
        }, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "error": f"BOM 对比失败：{exc}"}, ensure_ascii=False))
        sys.exit(1)


if __name__ == "__main__":
    main()
