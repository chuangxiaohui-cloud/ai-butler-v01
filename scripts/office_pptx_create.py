"""Create a simple project report .pptx from a JSON spec."""

import json
import sys

sys.stdout.reconfigure(encoding="utf-8")

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"python-pptx unavailable: {exc}"}))
    sys.exit(1)


def main() -> None:
    if len(sys.argv) < 3:
        print(json.dumps({"ok": False, "error": "usage: office_pptx_create.py <input.json> <output.pptx>"}))
        sys.exit(1)
    try:
        spec = json.load(open(sys.argv[1], encoding="utf-8"))
        accent = spec.get("accent")
        accent_color = None
        if accent and isinstance(accent, str) and len(accent) == 7 and accent.startswith("#"):
            accent_color = RGBColor.from_string(accent[1:])
        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = spec.get("title", "项目汇报")
        if accent_color:
            for run in title_slide.shapes.title.text_frame.paragraphs[0].runs:
                run.font.color.rgb = accent_color
        for slide_spec in spec.get("slides", []):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = slide_spec.get("title", "")
            if accent_color:
                for run in slide.shapes.title.text_frame.paragraphs[0].runs:
                    run.font.color.rgb = accent_color
            body = slide.placeholders[1].text_frame
            first = True
            for bullet in slide_spec.get("bullets", []):
                paragraph = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                paragraph.text = bullet
                paragraph.font.size = Pt(18)
                if accent_color and paragraph.runs:
                    paragraph.runs[0].font.color.rgb = accent_color
        prs.save(sys.argv[2])
        print(json.dumps({"ok": True, "path": sys.argv[2]}))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        sys.exit(1)


if __name__ == "__main__":
    main()
