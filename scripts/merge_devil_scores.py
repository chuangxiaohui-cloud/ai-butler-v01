#!/usr/bin/env python
"""Merge the owner's independent scores (xlsx) into bench/devil-v25/scores.json.

Usage:
  python scripts/merge_devil_scores.py

Reads:
  bench/devil-v25/scores.json            (current autoScore + initial score)
  bench/devil-v25/AI-Agent_魔鬼训练_v2.5_AI独立评分表.xlsx

Writes:
  bench/devil-v25/scores.json            (score=owner final, initialScore kept)
"""

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SCORES_PATH = ROOT / "bench" / "devil-v25" / "scores.json"
XLSX_PATH = (
    ROOT / "bench" / "devil-v25" / "AI-Agent_魔鬼训练_v2.5_AI独立评分表.xlsx"
)


def main() -> None:
    current = json.loads(SCORES_PATH.read_text(encoding="utf-8"))
    by_id = {s["id"]: s for s in current["scores"]}

    wb = load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    col = {name: i for i, name in enumerate(header)}
    merged = []
    for row in rows[1:]:
        sid = row[col["题号"]]
        owner_score = row[col["人工独立评分(0-3)"]]
        reason = row[col["评分理由"]]
        prev = by_id.get(sid, {})
        merged.append(
            {
                "id": sid,
                "score": int(owner_score),
                "initialScore": prev.get("score"),
                "autoScore": prev.get("autoScore"),
                "hardAnswer": False,
                "reason": reason or "",
            }
        )

    current["scores"] = merged
    current["note"] = (
        "2026-08-15 定稿：score 为 owner 人工独立评分（xlsx），"
        "initialScore 为 AI 初判分，autoScore 为脚本参考分，reason 为 owner 评分理由。"
    )
    SCORES_PATH.write_text(
        json.dumps(current, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"已合并 {len(merged)} 条到 {SCORES_PATH}")


if __name__ == "__main__":
    main()
