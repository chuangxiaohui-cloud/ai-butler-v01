"""Read an .xlsx first sheet into JSON rows."""

import json
import sys

sys.stdout.reconfigure(encoding="utf-8")

try:
    from openpyxl import load_workbook
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"openpyxl unavailable: {exc}"}))
    sys.exit(1)


def main() -> None:
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "usage: office_xlsx_read.py <input>"}))
        sys.exit(1)
    try:
        wb = load_workbook(sys.argv[1], data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell).strip() for cell in row])
            if len(rows) >= 2000:
                break
        wb.close()
        print(json.dumps({"ok": True, "headers": rows[0] if rows else [], "rows": rows}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        sys.exit(1)


if __name__ == "__main__":
    main()
