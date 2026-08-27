"""Read a .docx into plain text via python-docx (no Word COM)."""

import json
import sys

sys.stdout.reconfigure(encoding="utf-8")


def main() -> None:
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "usage: office_docx_read.py <input.docx>"}))
        sys.exit(1)
    try:
        import docx

        document = docx.Document(sys.argv[1])
        parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        print(json.dumps({
            "ok": True,
            "text": text,
            "chars": len(text),
            "lines": len(text.splitlines()) if text else 0,
        }, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "error": f"docx 读取失败（需要 python-docx）：{exc}"}, ensure_ascii=False))
        sys.exit(1)


if __name__ == "__main__":
    main()
